"""
The import template: public/import-template.xlsx.

    python3 scripts/build-import-template.py      (needs: pip install openpyxl)

Kept in the repository so the template can be rebuilt when a column changes,
rather than edited by hand in Excel and drifting from what the importer reads.
The one fact it must never get wrong is which name column is required: the
default locale (DEFAULT_LOCALE in src/lib/i18n/config.ts, and the locale with
is_default in the database) — which is English.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

CLAY = "B4593C"      # brand
CREAM = "FBF7F0"     # page
INK = "2C2523"

wb = Workbook()
ws = wb.active
ws.title = "Products"

COLUMNS = [
    ("slug", 22, "cizkejk-njujork"),
    ("position", 10, 0),
    ("price_rsd", 12, 570),
    ("whole_price_rsd", 17, 3420),
    ("tag", 12, "Hit"),
    ("formats", 20, "chilled"),
    ("weight_g", 11, 130),
    ("kcal", 9, 341),
    ("protein_g", 12, 5.8),
    ("fat_g", 9, 23.1),
    ("carbs_g", 11, 27.4),
    ("photo", 26, ""),
    ("published", 12, "yes"),
    ("name_sr", 26, "Čizkejk Njujork"),
    ("note_sr", 34, "Gust i kremast, na podlozi od keksa."),
    ("name_ru", 26, "Чизкейк Нью-Йорк"),
    ("note_ru", 34, "Плотный и сливочный, на песочной основе."),
    ("name_en", 26, "New York cheesecake"),
    ("note_en", 34, "Dense and creamy, on a biscuit base."),
]

REQUIRED = {"price_rsd", "formats", "name_en"}

head_fill = PatternFill("solid", fgColor=CLAY)
head_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="E6DED3")

for i, (name, width, _) in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=i, value=name + (" *" if name in REQUIRED else ""))
    cell.fill = head_fill
    cell.font = head_font
    cell.alignment = Alignment(vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = width

ws.row_dimensions[1].height = 26
ws.freeze_panes = "A2"

EXAMPLES = [
    ["cizkejk-njujork", 0, 570, 3420, "Hit", "chilled", 130, 341, 5.8, 23.1, 27.4, "", "yes",
     "Čizkejk Njujork", "Gust i kremast, na podlozi od keksa.",
     "Чизкейк Нью-Йорк", "Плотный и сливочный, на песочной основе.",
     "New York cheesecake", "Dense and creamy, on a biscuit base."],
    ["medovik", 1, 450, "", "", "chilled, frozen", 120, 362, 4.9, 19.7, 41.8, "", "yes",
     "Medovik", "Tanki medeni korovi sa kremom od pavlake.",
     "Медовик", "Тонкие медовые коржи со сметанным кремом.",
     "Medovik honey cake", "Thin honey layers with sour cream."],
    ["", 2, 320, "", "Novo", "ambient", 100, "", "", "", "", "", "no",
     "Brauni", "Gust, bez brašna, sa tamnom čokoladom.",
     "", "", "Brownie", "Dense, flourless, dark chocolate."],
]

for r, row in enumerate(EXAMPLES, start=2):
    for c, value in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=value if value != "" else None)
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(vertical="center", wrap_text=False)

# Storage words, offered rather than remembered.
# A hint, not a gate: several formats in one cell are normal ("chilled, frozen"),
# and Excel's list validation cannot express that — so it offers the three words
# and lets anything be typed.
storage = DataValidation(type="list", formula1='"chilled,frozen,ambient"', allow_blank=False)
storage.showErrorMessage = False
storage.prompt = "chilled · frozen · ambient. Several are fine: chilled, frozen"
storage.promptTitle = "Storage"
ws.add_data_validation(storage)
storage.add(f"F2:F400")

published = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
published.prompt = "Empty means yes."
published.promptTitle = "Live on the site"
ws.add_data_validation(published)
published.add("M2:M400")

# ── the second sheet: what each column is ──────────────────────────────────
notes = wb.create_sheet("How to fill it in")
notes.column_dimensions["A"].width = 26
notes.column_dimensions["B"].width = 14
notes.column_dimensions["C"].width = 86

title = notes.cell(row=1, column=1, value="One row per product")
title.font = Font(bold=True, size=14, color=INK)
notes.cell(row=2, column=1,
           value="Fill in the Products sheet, save the file, and upload it in the console: Products → Import a list.")
notes.cell(row=3, column=1,
           value="Rows are matched by slug: importing a corrected file updates what is already there instead of doubling it.")

for r, head in enumerate(["Column", "Required", "What it is"], start=1):
    cell = notes.cell(row=5, column=r, value=head)
    cell.fill = head_fill
    cell.font = head_font

DOCS = [
    ("slug", "no", "The product's address, in lowercase latin letters and hyphens. Leave it empty and it is made from the name — including from Cyrillic, so “Медовик” becomes medovik."),
    ("position", "no", "Lower comes first in the catalog. Leave it empty and the order of the rows is used."),
    ("price_rsd", "yes", "Price per piece, whole dinars. “1 250”, “1.250” and “1250 RSD” all read as 1250."),
    ("whole_price_rsd", "no", "The price of the whole cake. Filled in, the card offers the piece / whole-cake switch; empty, it does not."),
    ("tag", "no", "The word in the corner of the card — “Hit”, “Novo”. One language, shown exactly as typed. 24 characters."),
    ("formats", "yes", "How it keeps: chilled, frozen or ambient, several separated by commas. Serbian and Russian words are understood — “frizider”, “заморожено”, “sobna temperatura”."),
    ("weight_g", "no", "Grams. One piece — or the whole cake on a whole-cake product."),
    ("kcal, protein_g, fat_g, carbs_g", "no", "Per 100 g. All four or none: a declaration with a gap in it is a wrong one, so the card shows the panel only when it is whole."),
    ("photo", "no", "Usually left empty and uploaded in the console. A name already in the photo bucket works, and so does /products/<file>.jpg for photography already in the site."),
    ("published", "no", "yes or no. Empty means yes. A product with no name in the default language cannot be published, whatever this says."),
    ("name_sr, note_sr", "no", "Serbian. Leave a language out and it falls back to English."),
    ("name_ru, note_ru", "no", "Russian. Leave a language out and it falls back to English."),
    ("name_en, note_en", "yes", "English — the default language. Every other language falls back to it, so it is the one a product cannot go live without. The note is the single line under the name on the card."),
]

for i, (name, required, text) in enumerate(DOCS, start=6):
    notes.cell(row=i, column=1, value=name).font = Font(bold=True, color=INK)
    notes.cell(row=i, column=2, value=required)
    cell = notes.cell(row=i, column=3, value=text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    notes.row_dimensions[i].height = 30

tail = len(DOCS) + 7
notes.cell(row=tail, column=1, value="Saving from Excel, Numbers or LibreOffice").font = Font(bold=True, color=INK)
notes.cell(row=tail + 1, column=1,
           value="Any of .xlsx, .csv or .tsv works. The old .xls does not — open it and save as .xlsx.")
notes.cell(row=tail + 2, column=1,
           value="A row that cannot be read stops the whole import and comes back with its row number, so nothing lands half-done.")

import pathlib
wb.save(pathlib.Path(__file__).resolve().parent.parent / "public" / "import-template.xlsx")
print("written")
